import io
from datetime import datetime


def generate_excel(data: dict, export_type: str) -> bytes:
    """生成 Excel 文件"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "版本历史"

    # 标题行样式
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1A1A2E", end_color="1A1A2E", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    # 表头
    headers = ["版本号", "文件名", "文件大小", "标签", "变更说明", "创建时间"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # 数据行
    row = 2
    if export_type == "project":
        for series in data.get("series_list", []):
            for v in series.get("versions", []):
                ws.cell(row=row, column=1, value=f"v{v['version_number']}")
                ws.cell(row=row, column=2, value=v['filename'])
                ws.cell(row=row, column=3, value=format_size(v['file_size']))
                ws.cell(row=row, column=4, value=", ".join(v.get('tags', [])))
                ws.cell(row=row, column=5, value=v.get('change_note', ''))
                ws.cell(row=row, column=6, value=format_datetime(v['created_at']))
                row += 1
    else:  # series
        for v in data.get("versions", []):
            ws.cell(row=row, column=1, value=f"v{v['version_number']}")
            ws.cell(row=row, column=2, value=v['filename'])
            ws.cell(row=row, column=3, value=format_size(v['file_size']))
            ws.cell(row=row, column=4, value=", ".join(v.get('tags', [])))
            ws.cell(row=row, column=5, value=v.get('change_note', ''))
            ws.cell(row=row, column=6, value=format_datetime(v['created_at']))
            row += 1

    # 调整列宽
    for col in range(1, 7):
        ws.column_dimensions[chr(64 + col)].width = 15

    # 保存到字节流
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def generate_pdf(data: dict, export_type: str) -> bytes:
    """生成 PDF 文件"""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    # 注册中文字体（尝试使用系统字体）
    try:
        pdfmetrics.registerFont(TTFont('SimSun', 'C:/Windows/Fonts/simsun.ttc'))
        font_name = 'SimSun'
    except:
        font_name = 'Helvetica'

    # 创建 PDF
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, topMargin=20*mm, bottomMargin=20*mm)

    # 样式
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Title'],
        fontName=font_name,
        fontSize=18,
        spaceAfter=20*mm,
    )
    normal_style = ParagraphStyle(
        'CustomNormal',
        parent=styles['Normal'],
        fontName=font_name,
        fontSize=10,
    )

    # 内容
    elements = []

    # 标题
    if export_type == "project":
        title = f"项目版本历史 - {data.get('name', '')}"
    else:
        title = f"系列版本历史 - {data.get('drawing_number', '')}"
    elements.append(Paragraph(title, title_style))

    # 表格数据
    table_data = [["版本号", "文件名", "大小", "标签", "变更说明", "时间"]]

    if export_type == "project":
        for series in data.get("series_list", []):
            for v in series.get("versions", []):
                table_data.append([
                    f"v{v['version_number']}",
                    v['filename'][:20] + "..." if len(v['filename']) > 20 else v['filename'],
                    format_size(v['file_size']),
                    ", ".join(v.get('tags', []))[:15],
                    v.get('change_note', '')[:20],
                    format_datetime(v['created_at'])[:10],
                ])
    else:
        for v in data.get("versions", []):
            table_data.append([
                f"v{v['version_number']}",
                v['filename'][:20] + "..." if len(v['filename']) > 20 else v['filename'],
                format_size(v['file_size']),
                ", ".join(v.get('tags', []))[:15],
                v.get('change_note', '')[:20],
                format_datetime(v['created_at'])[:10],
            ])

    # 创建表格
    table = Table(table_data, colWidths=[15*mm, 40*mm, 20*mm, 25*mm, 35*mm, 25*mm])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1A1A2E')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), font_name),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('FONTNAME', (0, 1), (-1, -1), font_name),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('ALIGN', (0, 1), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F5F5F5')]),
        ('TOPPADDING', (0, 0), (-1, -1), 3*mm),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3*mm),
    ]))

    elements.append(table)
    doc.build(elements)

    return output.getvalue()


def format_size(bytes: int) -> str:
    if bytes < 1024:
        return f"{bytes} B"
    if bytes < 1024 * 1024:
        return f"{bytes / 1024:.1f} KB"
    return f"{bytes / 1024 / 1024:.1f} MB"


def format_datetime(dt) -> str:
    if isinstance(dt, str):
        return dt[:19]
    return dt.strftime("%Y-%m-%d %H:%M:%S")
